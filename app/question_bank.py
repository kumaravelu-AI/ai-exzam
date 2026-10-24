from __future__ import annotations
from pathlib import Path
from typing import Any
from openpyxl import load_workbook

from .config import KNOWLEDGEBASE_DIR

DEFAULT_COLUMNS = {
    "question_id": "question_id",
    "subject": "subject",
    "level": "level",
    "topic": "topic",
    "concept": "concept",
    "difficulty": "difficulty",
    "question_type": "question_type",
    "question_text": "question_text",
    "expected_answer": "expected_answer",
    "keywords": "keywords",
    "hint": "hint",
    "explanation": "explanation",
    "misconception_tag": "misconception_tag",
    "followup_question_id": "followup_question_id",
    "marks": "marks",
    "mode": "mode",
}

def _normalize(s: Any) -> str:
    return str(s).strip() if s is not None else ""

def load_questions(subject: str, level: str, topic: str) -> list[dict[str, Any]]:
    topic_dir = KNOWLEDGEBASE_DIR / subject / level / topic
    if not topic_dir.exists():
        return []
    out: list[dict[str, Any]] = []
    for xlsx in topic_dir.rglob("*.xlsx"):
        wb = load_workbook(str(xlsx), data_only=True)
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            headers = [_normalize(h).lower() for h in rows[0]]
            col_map = {name: idx for idx, name in enumerate(headers) if name}
            for row in rows[1:]:
                if row is None:
                    continue
                def val(key: str) -> str:
                    col = DEFAULT_COLUMNS[key]
                    idx = col_map.get(col)
                    if idx is None or idx >= len(row):
                        return ""
                    return _normalize(row[idx])

                q_text = val("question_text")
                if not q_text:
                    continue
                item = {
                    "question_id": val("question_id") or f"{topic}_{len(out)+1}",
                    "subject": val("subject") or subject,
                    "level": val("level") or level,
                    "topic": val("topic") or topic,
                    "concept": val("concept") or "general",
                    "difficulty": val("difficulty") or "medium",
                    "question_type": val("question_type") or "short_answer",
                    "question_text": q_text,
                    "expected_answer": val("expected_answer"),
                    "keywords": [k.strip() for k in val("keywords").split(",") if k.strip()],
                    "hint": val("hint"),
                    "explanation": val("explanation"),
                    "misconception_tag": val("misconception_tag"),
                    "followup_question_id": val("followup_question_id"),
                    "marks": int(val("marks") or "1"),
                    "mode": (val("mode") or "both").lower(),
                    "source_file": str(xlsx.relative_to(KNOWLEDGEBASE_DIR)),
                }
                out.append(item)
    return out

def question_matches_mode(question: dict[str, Any], mode: str) -> bool:
    qmode = question.get("mode", "both")
    return qmode in {"both", mode}

def get_questions_for_mode(subject: str, level: str, topic: str, mode: str) -> list[dict[str, Any]]:
    return [q for q in load_questions(subject, level, topic) if question_matches_mode(q, mode)]

def find_question_by_id(subject: str, level: str, topic: str, question_id: str) -> dict[str, Any] | None:
    for q in load_questions(subject, level, topic):
        if q.get("question_id") == question_id:
            return q
    return None
