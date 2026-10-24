from pathlib import Path
from typing import Iterable
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

def read_txt(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="ignore")

def read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    texts = []
    for page in reader.pages:
        try:
            texts.append(page.extract_text() or "")
        except Exception:
            texts.append("")
    return "\n".join(texts)

def read_docx(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text)

def read_xlsx(path: Path) -> str:
    wb = load_workbook(str(path), data_only=True)
    chunks: list[str] = []
    for ws in wb.worksheets:
        chunks.append(f"[SHEET] {ws.title}")
        rows = list(ws.iter_rows(values_only=True))
        for row in rows:
            vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
            if vals:
                chunks.append(" | ".join(vals))
    return "\n".join(chunks)

def supported_file(path: Path) -> bool:
    return path.suffix.lower() in {".txt", ".pdf", ".docx", ".xlsx"}

def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        return read_txt(path)
    if suffix == ".pdf":
        return read_pdf(path)
    if suffix == ".docx":
        return read_docx(path)
    if suffix == ".xlsx":
        return read_xlsx(path)
    return ""

def iter_supported_files(topic_dir: Path) -> Iterable[Path]:
    for path in topic_dir.rglob("*"):
        if path.is_file() and supported_file(path):
            yield path
